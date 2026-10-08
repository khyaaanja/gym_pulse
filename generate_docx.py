#!/usr/bin/env python3
"""Generate documentation.docx from DOCUMENTATION.md"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()

# Read the markdown file
with open('DOCUMENTATION.md', 'r') as f:
    content = f.read()

# Helper function to add styled heading
def add_heading(doc, text, level=1):
    if level == 1:
        doc.add_heading(text, level=0)
    elif level == 2:
        doc.add_heading(text, level=1)
    else:
        doc.add_heading(text, level=2)

# Helper to add styled paragraph
def add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph(text)
    for run in p.runs:
        if bold:
            run.bold = True
        if italic:
            run.italic = True
    return p

# Parse markdown and add to document
lines = content.split('\n')
i = 0
while i < len(lines):
    line = lines[i]

    # Skip empty lines at start
    if not line.strip():
        i += 1
        continue

    # Headings
    if line.startswith('# '):
        title = line[2:].strip()
        doc.add_heading(title, level=0)
        i += 1
    elif line.startswith('## '):
        heading = line[3:].strip()
        doc.add_heading(heading, level=1)
        i += 1
    elif line.startswith('### '):
        heading = line[4:].strip()
        doc.add_heading(heading, level=2)
        i += 1

    # Tables
    elif '|' in line and '---' in lines[i+1] if i+1 < len(lines) else False:
        # Parse markdown table
        header_row = [cell.strip() for cell in line.split('|')[1:-1]]
        rows = []
        i += 2  # Skip header and separator
        while i < len(lines) and '|' in lines[i] and not lines[i].startswith('#'):
            if lines[i].strip():
                row = [cell.strip() for cell in lines[i].split('|')[1:-1]]
                if len(row) == len(header_row):
                    rows.append(row)
            i += 1

        # Create table
        if header_row and rows:
            table = doc.add_table(rows=1, cols=len(header_row))
            table.style = 'Light Grid Accent 1'
            hdr_cells = table.rows[0].cells
            for j, header in enumerate(header_row):
                hdr_cells[j].text = header
                for paragraph in hdr_cells[j].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

            for row in rows:
                row_cells = table.add_row().cells
                for j, cell_text in enumerate(row):
                    row_cells[j].text = cell_text
        continue

    # Ordered/unordered lists
    elif line.strip().startswith(('- ', '* ', '+ ')):
        item = line.strip()[2:].strip()
        p = doc.add_paragraph(item, style='List Bullet')
        i += 1
    elif re.match(r'^\d+\.\s', line.strip()):
        item = re.sub(r'^\d+\.\s', '', line.strip())
        p = doc.add_paragraph(item, style='List Number')
        i += 1

    # Code blocks
    elif line.startswith('```'):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith('```'):
            code_lines.append(lines[i])
            i += 1
        code_block = '\n'.join(code_lines).strip()
        if code_block:
            p = doc.add_paragraph(code_block)
            p.style = 'Normal'
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
        i += 1  # Skip closing ```

    # Bold text patterns
    elif '**' in line:
        parts = line.split('**')
        p = doc.add_paragraph()
        for j, part in enumerate(parts):
            if j % 2 == 1:  # Bold part
                run = p.add_run(part)
                run.bold = True
            else:
                p.add_run(part)
        i += 1

    # Regular paragraphs
    elif line.strip():
        doc.add_paragraph(line.strip())
        i += 1
    else:
        i += 1

# Save document
doc.save('documentation.docx')
print("✓ Created documentation.docx")
